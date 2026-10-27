import io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header_row(ws, headers: list[str]):
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(header) + 4, 12)


def export_items(items: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    headers = ["Code", "Name", "Category", "Type", "Unit", "Standard Cost",
               "Reorder Point", "Lead Time (days)", "Active", "Description"]
    _style_header_row(ws, headers)
    for item in items:
        ws.append([
            item.get("code"), item.get("name"),
            item.get("category", {}).get("name") if item.get("category") else "",
            item.get("item_type"), item.get("unit"),
            float(item.get("standard_cost", 0)),
            float(item.get("reorder_point", 0)),
            item.get("lead_time_days", 0),
            "Yes" if item.get("is_active") else "No",
            item.get("description", ""),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_production_plans(plans: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Production Plans"
    headers = ["Plan No", "Item Code", "Item Name", "Planned Qty", "Unit",
               "Period Start", "Period End", "Status", "Actual Qty", "Achievement %", "Notes"]
    _style_header_row(ws, headers)
    for p in plans:
        planned = float(p.get("planned_qty", 0))
        actual = float(p.get("actual_qty_sum", 0))
        achievement = round(actual / planned * 100, 2) if planned else 0
        ws.append([
            p.get("plan_no"), p["item"]["code"], p["item"]["name"],
            planned, p.get("unit"),
            str(p.get("period_start")), str(p.get("period_end")),
            p.get("status"), actual, achievement, p.get("notes", ""),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_actual_outputs(outputs: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Actual Outputs"
    headers = ["Output No", "Item Code", "Item Name", "Output Date", "Actual Qty",
               "Good Qty", "Scrap Qty", "Unit", "Plan No", "Line", "Shift", "Operator", "Notes"]
    _style_header_row(ws, headers)
    for o in outputs:
        ws.append([
            o.get("output_no"), o["item"]["code"], o["item"]["name"],
            str(o.get("output_date")),
            float(o.get("actual_qty", 0)), float(o.get("good_qty", 0)),
            float(o.get("scrap_qty", 0)), o.get("unit"),
            o.get("plan_no", ""), o.get("line_id", ""),
            o.get("shift", ""), o.get("operator", ""), o.get("notes", ""),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_inventory_movements(movements: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory Movements"
    headers = ["Movement No", "Item Code", "Item Name", "Type", "Date",
               "Qty", "Unit", "From Location", "To Location", "Ref Type", "Ref ID", "Notes"]
    _style_header_row(ws, headers)
    for m in movements:
        ws.append([
            m.get("movement_no"), m["item"]["code"], m["item"]["name"],
            m.get("movement_type"), str(m.get("movement_date")),
            float(m.get("qty", 0)), m.get("unit"),
            m.get("from_location", ""), m.get("to_location", ""),
            m.get("reference_type", ""), m.get("reference_id", ""), m.get("notes", ""),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_items_xlsx(file_bytes: bytes) -> tuple[list[dict], list[dict]]:
    """Returns (valid_rows, error_rows). Expected columns: Code,Name,Type,Unit,Standard Cost,..."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    valid, errors = [], []

    for row_idx, row in enumerate(rows[1:], start=2):
        data = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
        if not data.get("code") or not data.get("name"):
            errors.append({"row": row_idx, "error": "Missing code or name", "data": data})
            continue
        valid.append({
            "code": str(data.get("code", "")).strip(),
            "name": str(data.get("name", "")).strip(),
            "item_type": str(data.get("type", "MATERIAL")).strip().upper(),
            "unit": str(data.get("unit", "EA")).strip(),
            "standard_cost": float(data.get("standard cost", 0) or 0),
            "reorder_point": float(data.get("reorder point", 0) or 0),
            "lead_time_days": int(data.get("lead time (days)", 0) or 0),
            "description": str(data.get("description", "") or ""),
        })
    return valid, errors
